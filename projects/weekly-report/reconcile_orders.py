"""Read-only reconciliation. Customer fields never leave the raw source files."""
from pathlib import Path
from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal
import argparse, csv, json, re, warnings
from datetime import timedelta
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
START, END = '', ''

def read(path):
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        wb = load_workbook(path, read_only=True, data_only=True)
    sh = wb.active
    sh.reset_dimensions()
    it = sh.iter_rows(values_only=True)
    headers = next(it)
    rows = [dict(zip(headers, row)) for row in it if any(v is not None for v in row)]
    wb.close()
    return rows

def money(v):
    if v in (None, '', '-', '--'): return Decimal(0)
    return Decimal(re.sub(r'[,元¥￥\s]', '', str(v)))

def count(v):
    if v in (None, '', '-', '--') or '指数' in str(v): return None
    return int(money(v))

def stamp(v):
    return v.isoformat(sep=' ') if isinstance(v, datetime) else str(v or '')

def run():
    global ROOT, START, END
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--manifest',type=Path,required=True)
    p.add_argument('--output',type=Path,required=True)
    args=p.parse_args()
    manifest=json.loads(args.manifest.read_text())
    START=manifest['period'][0]
    END=(datetime.fromisoformat(manifest['period'][1])+timedelta(days=1)).date().isoformat()
    for key in ('orders','aftersales','products'):
        path=Path(manifest[key]).expanduser()
        manifest[key]=str(path if path.is_absolute() else (args.manifest.parent/path).resolve())
    ROOT=args.output.resolve(); ROOT.mkdir(parents=True,exist_ok=True)
    orders = read(manifest['orders'])
    afters = read(manifest['aftersales'])
    products = read(manifest['products'])
    ids = [str(r['订单号']) for r in orders]
    assert len(ids)==len(set(ids)), 'Multi-line orders require line-level refund allocation'
    assert all(money(r['数量'])==1 for r in orders), 'Inspect multi-item quantities'
    assert all(START <= stamp(r['买家下单时间']) < END for r in orders)
    assert all(not r['买家支付时间'] or START<=stamp(r['买家支付时间'])<END for r in orders), 'Payment date boundary mismatch'
    byafter=defaultdict(list)
    for a in afters:
        if str(a['订单号']) in ids: byafter[str(a['订单号'])].append(a)
    pmap={}
    for p in products:
        sku=str(p['商品货号'])
        if sku in pmap: raise ValueError('Duplicate SPU SKU requires aggregation: '+sku)
        pmap[sku]=p
    buckets={}
    status=Counter()
    after_status=Counter()
    for o in orders:
        status[o['订单状态']]+=1
        if not o['买家支付时间']:continue
        sku=str(o['货号'])
        b=buckets.setdefault(sku,dict(sku=sku,name=o['商品名称'],raw=Decimal(0),closed=Decimal(0),other_refund=Decimal(0),active=Decimal(0),pending=Decimal(0),pending_requested=Decimal(0),paid_orders=0,closed_orders=0,active_orders=0,pending_orders=0,completed_refund_evidence=Decimal(0)))
        paid=money(o['出价金额（元）'])-money(o['卖家承担优惠金额（元）'])
        assert paid>=0
        b['raw']+=paid;b['paid_orders']+=1
        aa=byafter[str(o['订单号'])]
        succeeded=[a for a in aa if a['售后状态']=='售后成功']
        pending=[a for a in aa if a['售后状态'] not in ('售后成功','售后关闭')]
        for a in aa:after_status[a['售后状态']]+=1
        assert len(succeeded)<=1, 'Multiple successful refunds require manual allocation'
        refund=sum((money(a['退款金额（元）']) for a in succeeded),Decimal(0))
        b['completed_refund_evidence']+=refund
        if o['订单状态']=='交易关闭成功':
            b['closed']+=paid;b['closed_orders']+=1
            continue
        assert refund<=paid, 'Refund exceeds payment'
        retained=paid-refund;b['other_refund']+=refund;b['active']+=retained
        if retained>0:b['active_orders']+=1
        if pending and retained>0:
            b['pending']+=retained;b['pending_orders']+=1
            b['pending_requested']+=sum((money(a['退款金额（元）']) for a in pending),Decimal(0))
    output=[]
    for sku,b in buckets.items():
        assert sku in pmap, 'Missing SPU row: '+sku
        p=pmap[sku]
        assert b['raw']==money(p['支付金额']), 'SPU payment mismatch: '+sku
        assert b['paid_orders']==int(money(p['支付订单数'])), 'SPU order mismatch: '+sku
        uv=count(p['商详访问人数']);favorites=count(p['收藏用户数'])
        b.update(name=p['商品标题'],category=str(p['类目名称']).split('-')[0],detail_uv=uv,detail_pv=count(p['商详访问次数']),favorites=favorites,favorite_rate=round(favorites/uv*100,2) if uv and favorites is not None else None,conservative=b['active']-b['pending'])
        output.append({k:float(v) if isinstance(v,Decimal) else v for k,v in b.items()})
    output.sort(key=lambda r:(-r['active'],-r['raw'],r['sku']))
    sumkeys=['raw','closed','other_refund','active','pending','pending_requested','conservative','paid_orders','closed_orders','active_orders','pending_orders','completed_refund_evidence']
    total={k:round(sum(r[k] for r in output),2) for k in sumkeys}
    assert total['raw']==float(sum((money(p['支付金额']) for p in products),Decimal(0)))
    assert round(total['raw']-total['closed']-total['other_refund'],2)==total['active']
    total.update(all_order_rows=len(orders),paid_products=len(output),active_products=sum(r['active']>0 for r in output),product_pv=sum(count(p['商详访问次数']) or 0 for p in products),product_favorites=sum(count(p['收藏用户数']) or 0 for p in products))
    result=dict(period=manifest['period'],snapshot=manifest['snapshot'],summary=total,products=output,status_counts=dict(status),aftersales_status_counts=dict(after_status),sources=manifest,checks={'unique_order_ids':True,'order_and_payment_dates_within_week':True,'every_sku_payment_matches_spu':True,'every_sku_order_count_matches_spu':True,'sales_bridge_balances':True})
    (ROOT/'analysis.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
    print(json.dumps(total,ensure_ascii=False,indent=2))

if __name__=='__main__':run()
