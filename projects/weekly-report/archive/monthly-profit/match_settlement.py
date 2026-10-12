#!/usr/bin/env python3
"""Match September catalog order IDs with sales and refund settlement extracts."""

from collections import defaultdict
from decimal import Decimal
from pathlib import Path
import glob
import json
import warnings

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
DOWNLOADS = Path('/Users/nan/Downloads')


def read(path):
    with warnings.catch_warnings():
        warnings.filterwarnings('ignore', message='Workbook contains no default style')
        sheet = load_workbook(path, read_only=True, data_only=True).active
    sheet.reset_dimensions()
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)
    return [dict(zip(headers, row)) for row in rows if row[0]]


orders = json.loads((ROOT / 'analysis.json').read_text())['orders']
ids = {o['order_id'] for o in orders}
sales_path = Path(glob.glob(str(DOWNLOADS / '销售订单入账详情*.xlsx'))[0])
refund_path = Path(glob.glob(str(DOWNLOADS / '退货退款订单入账详情*.xlsx'))[0])
sales = {}
for r in read(sales_path):
    oid = str(r['订单号'])
    if oid in ids:
        if oid in sales:
            raise ValueError(f'duplicate sales settlement order {oid}')
        sales[oid] = {'amount_yuan': float(Decimal(str(r['应结金额（元）']))), 'status': str(r['入账状态'])}
refunds = defaultdict(list)
for r in read(refund_path):
    oid = str(r['订单号'])
    if oid in ids:
        refunds[oid].append({'amount_yuan': float(Decimal(str(r['应结金额（元）']))), 'status': str(r['入账状态'])})

result = {
    'sales_source': str(sales_path),
    'refund_source': str(refund_path),
    'sales_by_order': sales,
    'refund_by_order': dict(refunds),
}
(ROOT / 'settlement_match.json').write_text(json.dumps(result, ensure_ascii=False, indent=2))
active = [o for o in orders if o['valid_paid'] > 0]
shipped_closed = [o for o in orders if o['closed_paid'] > 0 and o['ship_time']]
matched_active = [o for o in active if o['order_id'] in sales]
matched_refunds = [o for o in shipped_closed if o['order_id'] in refunds]
print(json.dumps({
    'active': len(active), 'sales_matched': len(matched_active),
    'sales_matched_sum': round(sum(sales[o['order_id']]['amount_yuan'] for o in matched_active), 2),
    'sales_expected_sum': round(sum(o['expected_income'] for o in matched_active), 2),
    'shipped_closed': len(shipped_closed), 'refund_matched': len(matched_refunds),
    'confirmed_postage_loss': round(sum(o['postage_subsidy'] for o in matched_refunds), 2),
    'unconfirmed_postage_loss': round(sum(o['postage_subsidy'] for o in shipped_closed if o['order_id'] not in refunds), 2),
}, ensure_ascii=False, indent=2))
