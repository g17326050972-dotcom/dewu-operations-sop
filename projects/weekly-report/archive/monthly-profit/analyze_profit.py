#!/usr/bin/env python3
"""Analyze September DARA catalog orders without copying buyer information."""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
import json
from pathlib import Path
import warnings

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
CENT = Decimal("0.01")
GIFTBOX_WORDS = ("炽热挚爱", "钟意一生")


def money(value) -> Decimal:
    if value in (None, ""):
        return Decimal("0.00")
    return Decimal(str(value).replace(",", "").replace("元", "")).quantize(CENT)


def n(value: Decimal) -> float:
    return float(value.quantize(CENT, rounding=ROUND_HALF_UP))


def time(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value
    return datetime.fromisoformat(str(value))


def read(path: Path, fields: tuple[str, ...]):
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Workbook contains no default style")
        sheet = load_workbook(path, read_only=True, data_only=True).active
    sheet.reset_dimensions()
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)
    absent = sorted(set(fields) - set(headers))
    if absent:
        raise ValueError(f"{path.name}: missing fields {absent}")
    positions = {field: headers.index(field) for field in fields}
    return [
        {field: row[pos] if pos < len(row) else None for field, pos in positions.items()}
        for row in rows if any(value is not None for value in row)
    ]


ORDER_FIELDS = (
    "订单号", "商品名称", "货号", "SKU货号", "规格", "数量", "出价金额（元）",
    "卖家承担优惠金额（元）", "消费者邮费补贴金额（元）", "预计收入金额（元）",
    "订单状态", "买家下单时间", "买家支付时间", "卖家发货时间", "订单关闭时间",
)
AFTERSALE_FIELDS = (
    "售后单号", "订单号", "售后类型", "售后状态", "退款金额（元）", "买家申请售后时间",
)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--orders", type=Path, default=Path("/Users/nan/Downloads/得物订单导出 (2).xlsx"))
    parser.add_argument("--aftersales", type=Path, default=Path("/Users/nan/Downloads/售后单列表数据导出.xlsx"))
    parser.add_argument("--output", type=Path, default=ROOT / "analysis.json")
    args = parser.parse_args()

    catalog = json.loads((ROOT / "catalog_costs.json").read_text())["cost_by_sku"]
    orders = read(args.orders, ORDER_FIELDS)
    aftersales = read(args.aftersales, AFTERSALE_FIELDS)
    ids = [str(o["订单号"]) for o in orders]
    if len(ids) != len(set(ids)):
        raise ValueError("Duplicate order IDs; line-level allocation required")
    aftersale_ids = [str(a["售后单号"]) for a in aftersales]
    if len(aftersale_ids) != len(set(aftersale_ids)):
        raise ValueError("Duplicate after-sales IDs")

    after_by_order = defaultdict(list)
    for a in aftersales:
        after_by_order[str(a["订单号"])].append(a)

    status_counts = Counter()
    all_paid = Decimal(0)
    all_closed = Decimal(0)
    catalog_rows = []
    unmatched_paid_skus = set()
    for o in orders:
        sku = str(o["货号"] or "").strip().upper()
        status = str(o["订单状态"] or "").strip()
        status_counts[status] += 1
        qty = int(o["数量"] or 0)
        if qty != 1:
            raise ValueError(f"Unexpected quantity {qty} for order {o['订单号']}")
        offer = money(o["出价金额（元）"])
        discount = money(o["卖家承担优惠金额（元）"])
        paid = offer - discount if time(o["买家支付时间"]) else Decimal(0)
        if paid < 0:
            raise ValueError(f"Discount exceeds offer for order {o['订单号']}")
        all_paid += paid
        closed = status == "交易关闭成功"
        if closed:
            all_closed += paid
        if sku not in catalog:
            if paid:
                unmatched_paid_skus.add(sku)
            continue

        related = after_by_order[str(o["订单号"])]
        successful_refund = sum((money(a["退款金额（元）"]) for a in related if a["售后状态"] == "售后成功"), Decimal(0))
        pending = any(a["售后状态"] not in ("售后成功", "售后关闭") for a in related)
        valid_paid = paid if status not in ("交易关闭成功", "交易失败") else Decimal(0)
        # A closed order and its successful after-sales record describe the same event.
        partial_refund = successful_refund if valid_paid else Decimal(0)
        retained_paid = max(Decimal(0), valid_paid - partial_refund)
        gift = any(word in str(o["规格"] or "") for word in GIFTBOX_WORDS)
        catalog_rows.append({
            "order_id": str(o["订单号"]),
            "sku": sku,
            "sku_variant": str(o["SKU货号"] or ""),
            "name": str(o["商品名称"] or ""),
            "spec": str(o["规格"] or ""),
            "status": status,
            "order_time": str(o["买家下单时间"] or ""),
            "pay_time": str(o["买家支付时间"] or ""),
            "ship_time": str(o["卖家发货时间"] or ""),
            "offer": n(offer),
            "seller_discount": n(discount),
            "paid": n(paid),
            "postage_subsidy": n(money(o["消费者邮费补贴金额（元）"])),
            "expected_income": n(money(o["预计收入金额（元）"])),
            "closed_paid": n(paid if closed else Decimal(0)),
            "valid_paid": n(valid_paid),
            "successful_refund_on_active": n(partial_refund),
            "retained_paid": n(retained_paid),
            "pending_aftersale": pending,
            "successful_aftersale_count": sum(a["售后状态"] == "售后成功" for a in related),
            "unit_cost": catalog[sku]["cost_yuan"],
            "cost_source_mode": catalog[sku]["cost_source_mode"],
            "giftbox": gift,
        })

    result = {
        "basis": {
            "orders_file": str(args.orders), "aftersales_file": str(args.aftersales),
            "catalog_file": catalog[next(iter(catalog))]["sku_cell"].split("!")[0],
            "order_date_min": min(str(o["买家下单时间"]) for o in orders),
            "order_date_max": max(str(o["买家下单时间"]) for o in orders),
            "aftersales_file_modified": datetime.fromtimestamp(args.aftersales.stat().st_mtime).isoformat(timespec="seconds"),
        },
        "control": {
            "order_rows": len(orders), "order_status_counts": dict(status_counts),
            "all_paid": n(all_paid), "all_closed_paid": n(all_closed),
            "all_current_effective_paid": n(all_paid - all_closed),
            "catalog_skus": len(catalog), "catalog_order_rows": len(catalog_rows),
            "catalog_distinct_order_skus": sorted({r["sku"] for r in catalog_rows}),
            "unmatched_paid_skus": sorted(unmatched_paid_skus),
        },
        "orders": catalog_rows,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2))
    print(json.dumps({"basis": result["basis"], "control": {k: v for k, v in result["control"].items() if k != "unmatched_paid_skus"}}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
