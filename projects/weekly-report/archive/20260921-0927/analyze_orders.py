#!/usr/bin/env python3
"""Reconcile DARA's weekly order, after-sales and product exports.

Only aggregate product and status fields are written. Buyer, address, phone,
logistics, raw order numbers and after-sales numbers are never exported.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal
from pathlib import Path
import json
import warnings

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
ORDERS = Path('/Users/nan/Downloads/得物订单导出 (1).xlsx')
AFTERSALES = Path('/Users/nan/Downloads/售后单列表数据导出.xlsx')
PRODUCTS = Path('/Users/nan/Downloads/商品（SPU）明细导出.xlsx')
OUTPUT = ROOT / 'order_analysis.json'
PERIOD_START = datetime(2026, 9, 21)
PERIOD_END = datetime(2026, 9, 28)
CENT = Decimal('0.01')

ORDER_FIELDS = (
    '订单号', '商品名称', '货号', 'SKU货号', '数量', '出价金额（元）',
    '卖家承担优惠金额（元）', '预计收入金额（元）', '订单状态', '关闭原因',
    '买家下单时间', '买家支付时间', '订单关闭时间',
)
AFTERSALE_FIELDS = (
    '售后单号', '订单号', '商品货号', '售后类型', '售后状态',
    '价格（元）', '卖家承担优惠金额（元）', '退款金额（元）',
    '买家申请售后时间',
)
PRODUCT_FIELDS = (
    '商品货号', '商品标题', '类目名称', '支付金额', '支付订单数',
    '退货金额（支付时间）', '退货订单数（支付时间）',
    '退货金额（退款时间）', '退货订单数（退款时间）',
)


def records(path: Path, fields: tuple[str, ...]) -> list[dict]:
    with warnings.catch_warnings():
        warnings.filterwarnings('ignore', message='Workbook contains no default style')
        sheet = load_workbook(path, read_only=True, data_only=True).active
    # Dewu's exported XML contains rows beyond its incorrect A1 dimension.
    sheet.reset_dimensions()
    source = sheet.iter_rows(values_only=True)
    headers = next(source)
    missing = set(fields) - set(headers)
    if missing:
        raise ValueError(f'{path.name}: missing columns {sorted(missing)}')
    positions = {field: headers.index(field) for field in fields}
    result = []
    for row in source:
        if not any(value is not None for value in row):
            continue
        # Never retain address, contact, logistics or buyer fields.
        result.append({field: row[index] if index < len(row) else None
                       for field, index in positions.items()})
    return result


def amount(value) -> Decimal:
    if value is None or value == '':
        return Decimal('0')
    return Decimal(str(value).replace(',', '').replace('元', '')).quantize(CENT)


def yuan(value: Decimal) -> float:
    return float(value.quantize(CENT))


def instant(value) -> datetime | None:
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value
    return datetime.fromisoformat(str(value))


def product_bucket() -> dict:
    return {
        'name': '', 'category': '', 'rawPaidAmount': Decimal('0'),
        'closedAmount': Decimal('0'), 'completedRefundAmount': Decimal('0'),
        'actualSalesReference': Decimal('0'), 'pendingAfterSalesAmount': Decimal('0'),
        'pendingAfterSalesRequestedRefund': Decimal('0'),
        'conservativeRetainedReference': Decimal('0'),
        'productReportedRefundByPaymentTime': Decimal('0'),
        'productReportedRefundByRefundTime': Decimal('0'),
        'rawPaidOrders': 0, 'closedOrders': 0, 'completedRefundOrders': 0,
        'activeOrders': 0, 'pendingAfterSalesOrders': 0,
        'conservativeRetainedOrders': 0,
    }


def main() -> None:
    orders = records(ORDERS, ORDER_FIELDS)
    aftersales = records(AFTERSALES, AFTERSALE_FIELDS)
    products = records(PRODUCTS, PRODUCT_FIELDS)
    if len(orders) != 185:
        raise ValueError(f'Expected 185 exported order lines, got {len(orders)}')

    order_ids = [str(row['订单号']) for row in orders]
    if len(set(order_ids)) != len(order_ids):
        # This week's export has one line per order. Future multi-line exports
        # need a line-level key before order-level after-sales can be assigned.
        raise ValueError('Order IDs repeat; inspect multi-item allocation first')
    order_by_id = dict(zip(order_ids, orders))
    out_of_period = sum(not (PERIOD_START <= instant(row['买家下单时间']) < PERIOD_END)
                        for row in orders)
    if out_of_period:
        raise ValueError(f'{out_of_period} order lines outside 2026-09-21..27')
    paid_out_of_period = sum(
        bool(instant(row['买家支付时间'])) and
        not (PERIOD_START <= instant(row['买家支付时间']) < PERIOD_END)
        for row in orders
    )

    aftersale_ids = [str(row['售后单号']) for row in aftersales]
    if len(set(aftersale_ids)) != len(aftersale_ids):
        raise ValueError('Duplicate after-sales IDs; investigate before summing')
    matched = [row for row in aftersales if str(row['订单号']) in order_by_id]
    matched_by_order: dict[str, list[dict]] = defaultdict(list)
    for row in matched:
        matched_by_order[str(row['订单号'])].append(row)
    if any(len(rows) > 1 for rows in matched_by_order.values()):
        raise ValueError('Multiple after-sales records for an order need manual allocation')

    product_export: dict[str, dict] = defaultdict(product_bucket)
    for row in products:
        sku = str(row['商品货号'] or '').strip()
        bucket = product_export[sku]
        bucket['name'] = str(row['商品标题'] or bucket['name'])
        bucket['category'] = str(row['类目名称'] or bucket['category'])
        bucket['rawPaidAmount'] += amount(row['支付金额'])
        bucket['rawPaidOrders'] += int(amount(row['支付订单数']))
        bucket['productReportedRefundByPaymentTime'] += amount(row['退货金额（支付时间）'])
        bucket['productReportedRefundByRefundTime'] += amount(row['退货金额（退款时间）'])

    by_sku: dict[str, dict] = defaultdict(product_bucket)
    status_counts = Counter()
    status_paid = defaultdict(lambda: Decimal('0'))
    aftersale_status_counts = Counter()
    aftersale_status_refunds = defaultdict(lambda: Decimal('0'))
    line_issues = Counter()
    for order_id, row in order_by_id.items():
        sku = str(row['货号'] or '').strip()
        if not sku:
            line_issues['missingSku'] += 1
        if sku != str(row['SKU货号'] or '').strip():
            line_issues['skuVsSkuCodeMismatch'] += 1
        if amount(row['数量']) != Decimal('1'):
            line_issues['quantityNotOne'] += 1
        status = str(row['订单状态'] or '').strip()
        status_counts[status] += 1
        bid = amount(row['出价金额（元）'])
        seller_discount = amount(row['卖家承担优惠金额（元）'])
        paid = bid - seller_discount if instant(row['买家支付时间']) else Decimal('0')
        if paid < 0:
            line_issues['negativePaid'] += 1
        status_paid[status] += paid

        bucket = by_sku[sku]
        bucket['name'] = product_export[sku]['name'] or str(row['商品名称'] or '')
        bucket['category'] = product_export[sku]['category']
        if not instant(row['买家支付时间']):
            if status != '交易失败':
                line_issues['unpaidUnexpectedStatus'] += 1
            continue
        bucket['rawPaidAmount'] += paid
        bucket['rawPaidOrders'] += 1

        matched_after = matched_by_order.get(order_id, [])
        completed_refund = Decimal('0')
        pending_refund = Decimal('0')
        pending = False
        for after in matched_after:
            after_sku = str(after['商品货号'] or '').strip()
            if after_sku != sku:
                line_issues['aftersaleSkuMismatch'] += 1
            after_status = str(after['售后状态'] or '').strip()
            refund = amount(after['退款金额（元）'])
            aftersale_status_counts[after_status] += 1
            aftersale_status_refunds[after_status] += refund
            if after_status == '售后成功':
                completed_refund += refund
                bucket['completedRefundOrders'] += 1
            elif after_status not in ('售后关闭',):
                pending_refund += refund
                pending = True
        bucket['completedRefundAmount'] += completed_refund
        if completed_refund > paid:
            line_issues['refundExceedsPaid'] += 1

        if status == '交易关闭成功':
            # Full product order closed. Successful refund can be below paid
            # because of postage/other deductions; do not treat that difference
            # as a retained product sale or subtract the refund a second time.
            bucket['closedAmount'] += paid
            bucket['closedOrders'] += 1
            if not matched_after or any(a['售后状态'] != '售后成功' for a in matched_after):
                line_issues['closedWithoutSuccessfulAftersale'] += 1
            continue
        if status not in ('交易成功', '待买家收货', '待物流揽收'):
            line_issues['paidUnexpectedStatus'] += 1
        if completed_refund:
            line_issues['successfulRefundOnNonclosedOrder'] += 1
        retained = max(Decimal('0'), paid - completed_refund)
        bucket['actualSalesReference'] += retained
        bucket['activeOrders'] += 1
        if pending:
            # Conservative scenario excludes the whole still-returning order.
            bucket['pendingAfterSalesAmount'] += retained
            bucket['pendingAfterSalesRequestedRefund'] += pending_refund
            bucket['pendingAfterSalesOrders'] += 1
        else:
            bucket['conservativeRetainedReference'] += retained
            bucket['conservativeRetainedOrders'] += 1

    products_out = []
    money_keys = (
        'rawPaidAmount', 'closedAmount', 'completedRefundAmount',
        'actualSalesReference', 'pendingAfterSalesAmount',
        'pendingAfterSalesRequestedRefund', 'conservativeRetainedReference',
        'productReportedRefundByPaymentTime',
        'productReportedRefundByRefundTime',
    )
    for sku, bucket in by_sku.items():
        bucket['productReportedRefundByPaymentTime'] = product_export[sku]['productReportedRefundByPaymentTime']
        bucket['productReportedRefundByRefundTime'] = product_export[sku]['productReportedRefundByRefundTime']
        products_out.append({'sku': sku, **{
            key: yuan(value) if key in money_keys else value
            for key, value in bucket.items()
        }})
    products_out.sort(key=lambda x: (-x['actualSalesReference'],
                                     -x['conservativeRetainedReference'], x['sku']))

    def total(field: str) -> Decimal | int:
        return sum((b[field] for b in by_sku.values()),
                   Decimal('0') if field in money_keys else 0)

    summary = {
        'allOrderRows': len(orders),
        'rawPaidOrders': total('rawPaidOrders'),
        'rawPaidAmount': yuan(total('rawPaidAmount')),
        'closedOrders': total('closedOrders'),
        'closedAmount': yuan(total('closedAmount')),
        'completedRefundOrders': total('completedRefundOrders'),
        'completedRefundAmount': yuan(total('completedRefundAmount')),
        'closedPaidLessCompletedRefund': yuan(total('closedAmount') - total('completedRefundAmount')),
        'activeOrders': total('activeOrders'),
        'actualSalesReference': yuan(total('actualSalesReference')),
        'pendingAfterSalesOrders': total('pendingAfterSalesOrders'),
        'pendingAfterSalesAmount': yuan(total('pendingAfterSalesAmount')),
        'pendingAfterSalesRequestedRefund': yuan(total('pendingAfterSalesRequestedRefund')),
        'conservativeRetainedOrders': total('conservativeRetainedOrders'),
        'conservativeRetainedReference': yuan(total('conservativeRetainedReference')),
        'productReportedRefundByPaymentTime': yuan(sum((b['productReportedRefundByPaymentTime'] for b in product_export.values()), Decimal('0'))),
        'productReportedRefundByRefundTime': yuan(sum((b['productReportedRefundByRefundTime'] for b in product_export.values()), Decimal('0'))),
        'statusBreakdown': [
            {'status': status, 'orderRows': count, 'rawPaidAmount': yuan(status_paid[status])}
            for status, count in sorted(status_counts.items())
        ],
        'matchedAfterSalesRows': len(matched),
        'matchedAfterSalesStatusBreakdown': [
            {'status': status, 'rows': count, 'listedRefundAmount': yuan(aftersale_status_refunds[status])}
            for status, count in sorted(aftersale_status_counts.items())
        ],
    }

    gross_mismatch = [sku for sku, bucket in by_sku.items()
                      if bucket['rawPaidAmount'] != product_export[sku]['rawPaidAmount']]
    count_mismatch = [sku for sku, bucket in by_sku.items()
                      if bucket['rawPaidOrders'] != product_export[sku]['rawPaidOrders']]
    checks = {
        'orderIdUnique': len(set(order_ids)) == len(order_ids),
        'aftersaleIdUnique': len(set(aftersale_ids)) == len(aftersale_ids),
        'allOrderRowsWithinWeek': out_of_period == 0,
        'allPaidTimesWithinWeek': paid_out_of_period == 0,
        'allPaidSkuAmountsMatchProductExport': not gross_mismatch,
        'allPaidSkuOrderCountsMatchProductExport': not count_mismatch,
        'grossAmountMismatchSkuCount': len(gross_mismatch),
        'orderCountMismatchSkuCount': len(count_mismatch),
        'rawLessClosedEqualsActual': (total('rawPaidAmount') - total('closedAmount') == total('actualSalesReference')),
        'actualLessPendingEqualsConservative': (total('actualSalesReference') - total('pendingAfterSalesAmount') == total('conservativeRetainedReference')),
        'closedAndCompletedAftersalesSameCount': total('closedOrders') == total('completedRefundOrders'),
        'lineIssues': dict(line_issues),
    }
    result = {
        'period': {'start': '2026-09-21', 'endInclusive': '2026-09-27',
                   'basis': '买家下单时间；全部 185 条恰好也在本周支付或未支付'},
        'sources': {
            'orders': {'filename': ORDERS.name, 'rows': len(orders), 'fileModified': datetime.fromtimestamp(ORDERS.stat().st_mtime).isoformat(timespec='seconds')},
            'aftersales': {'filename': AFTERSALES.name, 'rows': len(aftersales), 'fileModified': datetime.fromtimestamp(AFTERSALES.stat().st_mtime).isoformat(timespec='seconds')},
            'products': {'filename': PRODUCTS.name, 'rows': len(products), 'fileModified': datetime.fromtimestamp(PRODUCTS.stat().st_mtime).isoformat(timespec='seconds')},
        },
        'summary': summary,
        'products': products_out,
        'checks': checks,
        'methodology': [
            '订单支付额 = 出价金额 − 卖家承担优惠金额；按货号及总额与商品 SPU 导出的支付金额、支付订单数逐项一致。预计收入不是买家支付金额。',
            '已支付后“交易关闭成功”的商品订单整笔从销售额剔除；其“售后成功”的退款金额只是佐证，可能因运费等少于原支付额，不再二次扣除。',
            '实际销售参考 = 原始支付 − 已关闭订单；其中仍有进行中的退货。稳健留存参考进一步剔除整笔待处理退货订单，属于谨慎情景，不是最终结算额。',
            '按订单号连接售后单；本周一单一商品、数量均为 1，故不存在多商品金额分摊。',
        ],
        'limitations': [
            '订单导出不含买家实付或独立退款成功金额；支付额由两列重建并经商品导出逐货号核对。',
            '进行中的 11 笔退货最终可能成功、拒绝或部分退款；当前不能给出最终净销售额。',
            '商品导出的按支付时间退货金额与售后成功、关闭状态存在时间/定义差异；不能在本口径上直接再次扣减。',
            '导出只涵盖截至文件生成时可见的订单和售后状态；后续关闭或退款会改变留存额。',
        ],
    }
    OUTPUT.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    print(f'Wrote {OUTPUT} with {len(products_out)} paid SKUs; summary: {summary}')
    if not all(value for key, value in checks.items() if isinstance(value, bool)) or line_issues:
        raise ValueError(f'Validation issue: {checks}')


if __name__ == '__main__':
    main()
