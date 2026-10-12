"""Summarize the 2026-09-21..27 SPU export without changing the source workbook."""

from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal
from hashlib import sha256
from pathlib import Path
import json
import re
import warnings

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
WORK = Path(__file__).resolve().parent
SOURCE = Path('/Users/nan/Downloads/商品（SPU）明细导出.xlsx')
OUTPUT = WORK / 'product_analysis.json'

MONEY_COLS = [
    '支付金额', '退货金额（退款时间）', '退货金额（支付时间）',
]
COUNT_COLS = [
    '支付订单数', '商详访问人数', '商详访问次数', '收藏用户数',
    '退货订单数（退款时间）', '退货订单数（支付时间）', '支付用户数',
]
METRIC_COLS = MONEY_COLS + COUNT_COLS


def decimal_value(value):
    if value is None or value == '-':
        return Decimal('0')
    if '指数' in str(value):
        return Decimal('0')  # A platform index is not an absolute visitor count.
    cleaned = re.sub(r'[\s,元]', '', str(value))
    if not cleaned:
        return Decimal('0')
    return Decimal(cleaned)


def number(value):
    return float(value) if value % 1 else int(value)


def image_for(sku):
    for folder in [WORK / 'images', ROOT / 'work/20260907-0913/images', ROOT / 'work/dara-site/public/top_images']:
        for extension in ['.jpeg', '.jpg', '.png', '.webp']:
            path = folder / (sku + extension)
            if path.is_file() and path.stat().st_size > 2000:
                return str(path.relative_to(ROOT))
    return None


with warnings.catch_warnings():
    warnings.filterwarnings('ignore', message='Workbook contains no default style')
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
sheet = workbook.active
declared_dimension = sheet.calculate_dimension()
sheet.reset_dimensions()  # This export declares A1:A1 despite containing many columns.
iterator = sheet.iter_rows(values_only=True)
headers = list(next(iterator))
rows = [dict(zip(headers, row)) for row in iterator if any(value is not None for value in row)]
workbook.close()
indexed_visit_cells = {
    column: sum('指数' in str(row.get(column)) for row in rows)
    for column in ['商详访问人数', '商详访问次数']
}

required = {'SPUID', '商品标题', '商品货号', '类目名称', '支付金额', '支付订单数', '商详访问人数', '收藏用户数'}
missing = required - set(headers)
if missing:
    raise ValueError(f'Missing required columns: {sorted(missing)}')

for row in rows:
    row['_metric'] = {column: decimal_value(row[column]) for column in METRIC_COLS}

totals = {column: sum((row['_metric'][column] for row in rows), Decimal('0')) for column in METRIC_COLS}
sku_counts = Counter(str(row['商品货号']) for row in rows)
spu_counts = Counter(str(row['SPUID']) for row in rows)
category_sums = defaultdict(lambda: defaultdict(lambda: Decimal('0')))
category_active = Counter()

def product(row):
    m = row['_metric']
    visitors = m['商详访问人数']
    favorites = m['收藏用户数']
    sku = str(row['商品货号'])
    category = str(row['类目名称']).split('-')[0]
    cvr_raw = row.get('商详转化率')
    cvr = float(str(cvr_raw).replace('%', '')) if cvr_raw not in [None, '', '-'] else None
    return {
        'spuId': str(row['SPUID']),
        'sku': sku,
        'name': str(row['商品标题']),
        'category': category,
        'paidAmount': number(m['支付金额']),
        'paidOrders': number(m['支付订单数']),
        'paidUsers': number(m['支付用户数']),
        'detailUV': number(visitors),
        'detailPV': number(m['商详访问次数']),
        'favorites': number(favorites),
        'favoriteRatePct': round(float(favorites / visitors * 100), 2) if visitors else None,
        'backendConversionRatePct': cvr,
        'returnAmountByRefundTime': number(m['退货金额（退款时间）']),
        'returnOrdersByRefundTime': number(m['退货订单数（退款时间）']),
        'returnAmountByPaymentTime': number(m['退货金额（支付时间）']),
        'returnOrdersByPaymentTime': number(m['退货订单数（支付时间）']),
        'cachedImage': image_for(sku),
    }


for row in rows:
    category = str(row['类目名称']).split('-')[0]
    for column in METRIC_COLS:
        category_sums[category][column] += row['_metric'][column]
    if row['_metric']['支付金额'] > 0:
        category_active[category] += 1

paid_products = [product(row) for row in rows if row['_metric']['支付金额'] > 0]
paid_products.sort(key=lambda item: (-item['paidAmount'], -item['paidOrders'], item['sku']))
refund_time_products = [product(row) for row in rows if row['_metric']['退货金额（退款时间）'] > 0]
refund_time_products.sort(key=lambda item: (-item['returnAmountByRefundTime'], item['sku']))
top10 = paid_products[:10]
top10_missing = [item['sku'] for item in top10 if not item['cachedImage']]

category_rows = []
for category, m in sorted(category_sums.items(), key=lambda item: -item[1]['支付金额']):
    category_rows.append({
        'category': category,
        'paidAmount': number(m['支付金额']),
        'paidOrders': number(m['支付订单数']),
        'activeProducts': category_active[category],
        'paidSharePct': round(float(m['支付金额'] / totals['支付金额'] * 100), 2) if totals['支付金额'] else None,
        'returnAmountByPaymentTime': number(m['退货金额（支付时间）']),
        'returnAmountByRefundTime': number(m['退货金额（退款时间）']),
    })

report = {
    'source': {
        'path': str(SOURCE),
        'sha256': sha256(SOURCE.read_bytes()).hexdigest(),
        'sheet': sheet.title,
        'declaredDimension': declared_dimension,
        'periodFromTask': {'start': '2026-09-21', 'end': '2026-09-27'},
        'periodPresentInWorkbook': False,
        'fileModifiedAt': datetime.fromtimestamp(SOURCE.stat().st_mtime).isoformat(timespec='seconds'),
        'headers': headers,
    },
    'aggregate': {
        'spuRows': len(rows),
        'uniqueSpuIds': len(spu_counts),
        'uniqueSkus': len(sku_counts),
        'duplicateSkus': [sku for sku, count in sku_counts.items() if count > 1],
        'activePaidProducts': len(paid_products),
        'rawPaidAmount': number(totals['支付金额']),
        'rawPaidOrders': number(totals['支付订单数']),
        'rawAverageOrderAmount': round(float(totals['支付金额'] / totals['支付订单数']), 2) if totals['支付订单数'] else None,
        'sumProductDetailUV': number(totals['商详访问人数']),
        'sumProductDetailPV': number(totals['商详访问次数']),
        'indexLabeledVisitCellsExcludedFromSums': indexed_visit_cells,
        'sumProductFavorites': number(totals['收藏用户数']),
        'returnAmountByPaymentTime': number(totals['退货金额（支付时间）']),
        'returnOrdersByPaymentTime': number(totals['退货订单数（支付时间）']),
        'returnAmountByRefundTime': number(totals['退货金额（退款时间）']),
        'returnOrdersByRefundTime': number(totals['退货订单数（退款时间）']),
        'rawTop10PaidAmount': sum(item['paidAmount'] for item in top10),
        'rawTop10PaidOrders': sum(item['paidOrders'] for item in top10),
    },
    'categories': category_rows,
    'topPaidProducts': top10,
    'paidProducts': paid_products,
    'refundTimeProducts': refund_time_products,
    'imageCoverage': {
        'top10Cached': len(top10) - len(top10_missing),
        'top10MissingSkus': top10_missing,
        'top10ImagePaths': {item['sku']: item['cachedImage'] for item in top10},
    },
    'definitionsAndLimits': [
        '商品支付金额是后台 SPU 原始支付口径；退货或交易关闭的订单仍可能留在支付金额内，不能把本文件金额当作实际销售。',
        '退货金额（支付时间）按支付时间归属，退货金额（退款时间）按退款完成时间归属；两列不能同时从支付金额扣除。',
        '按支付时间退货金额也不足以完成逐单实销核算。需用订单导出核对交易关闭、部分退款、退款状态和订单行 SKU。',
        '商详访问人数与收藏用户数在此为逐商品加总，跨商品可能重复，不能当作全店去重 UV 或收藏人数。',
        '标注“指数”的商详访问字段不是绝对人数或次数，已从逐商品访问量加总中排除。',
        '单款收藏率 = 该款收藏用户数 / 该款商详访问人数。后台商详转化率沿用源文件，不另造全店转化率。',
        '图片只记录工作区现有缓存，缺图商品应从对应 SKU 商品页取得真实原图，不能使用占位图。',
        'Excel 不包含日期范围列；2026-09-21 至 09-27 的范围来自任务说明，仍需与后台筛选状态核对。',
    ],
}

OUTPUT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
print(f'Wrote {OUTPUT}; {len(rows)} SPUs, paid {number(totals["支付金额"])} yuan / {number(totals["支付订单数"])} orders, top10 image cache {len(top10)-len(top10_missing)}/10')
