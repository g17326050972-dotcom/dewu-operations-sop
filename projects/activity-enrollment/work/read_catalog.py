import json, warnings, collections, pathlib, re
import openpyxl

src=pathlib.Path('/Users/nan/Desktop/戴拉银饰货盘.xlsx')
with warnings.catch_warnings(record=True) as notices:
    warnings.simplefilter('always')
    wb=openpyxl.load_workbook(src,read_only=True,data_only=False,keep_links=True)
    cache=openpyxl.load_workbook(src,read_only=True,data_only=True,keep_links=True)
out={'source':str(src),'warnings':sorted(set(str(x.message) for x in notices)),'sheets':[],'defined_names':[{'name':n.name,'value':n.attr_text} for n in wb.defined_names.values()]}
for ws in wb.worksheets:
    cs=cache[ws.title]
    cached_cells={cell.coordinate:cell.value for cached_row in cs.iter_rows() for cell in cached_row if cell.value is not None}
    rows=[]; formulas=[]; comments=[]
    for row_number,row in enumerate(ws.iter_rows(),1):
        cells={}
        for cell in row:
            if cell.value is not None:
                value=cell.value
                if cell.data_type=='f':
                    formula=value if isinstance(value,str) else value.text
                    item={'cell':cell.coordinate,'formula':formula,'cached':cached_cells.get(cell.coordinate)}
                    formulas.append(item);cells[cell.column_letter]=item
                else:cells[cell.column_letter]=value
            if getattr(cell,'comment',None):comments.append({'cell':cell.coordinate,'text':cell.comment.text})
        if cells:rows.append({'row':row_number,'cells':cells})
    out['sheets'].append({'name':ws.title,'state':ws.sheet_state,'dimensions':ws.calculate_dimension(),'rows':rows,'formulas':formulas,'comments':comments})
pathlib.Path('work').mkdir(exist_ok=True)
pathlib.Path('work/catalog_raw.json').write_text(json.dumps(out,ensure_ascii=False,indent=2,default=str))
sheet=out['sheets'][0]
records=[]; missing=[]; group=collections.defaultdict(list); discrepancies=[]
def val(cells,col):
    v=cells.get(col)
    return v.get('cached') if isinstance(v,dict) else v
for row in sheet['rows'][1:]:
    n=row['row']; cells=row['cells']
    for col,v in cells.items():
        if isinstance(v,dict) and col!='C':
            formula=re.sub(r'(?<=[A-Z])\$?\d+','{row}',v['formula'])
            group[(col,formula)].append(n)
    price=val(cells,'D');discount=val(cells,'E');cost=val(cells,'G');p=val(cells,'R')
    if not isinstance(price,(int,float)) or not isinstance(cost,(int,float)) or not isinstance(discount,(int,float)):
        missing.append(n)
    rec={'row':n,'product_id':str(val(cells,'A')),'article_no':val(cells,'B'),'bid_price':price,'configured_discount':discount,'current_take_home_price':val(cells,'F'),'product_cost':cost,'platform_fee_current':val(cells,'H'),'packaging':val(cells,'I'),'materials':val(cells,'J'),'freight':val(cells,'K'),'operation_fee_current':val(cells,'L'),'tax_current':val(cells,'M'),'gross_profit_current':val(cells,'N'),'gross_margin_current':val(cells,'O'),'break_even_price_cached':val(cells,'P'),'activity_discount_cached':val(cells,'Q'),'activity_price_cached':p,'activity_profit_cached':val(cells,'S'),'activity_margin_cached':val(cells,'T'),'formulas':{k:v['formula'] for k,v in cells.items() if isinstance(v,dict) and k!='C'}}
    if isinstance(cost,(int,float)) and isinstance(price,(int,float)) and price>0:
        threshold=max((cost+59.5)/.87,(cost+14.5)/.65)
        rec['activity_price_at_3pct_margin']=threshold
        rec['discount_at_3pct_margin']=threshold/price*10
        configured_price=price*discount/10
        configured_profit=configured_price-max(configured_price*.22,45)-14.5-configured_price*.10-cost
        rec['activity_price_at_configured_discount']=configured_price
        rec['activity_profit_at_configured_discount']=configured_profit
        rec['activity_margin_at_configured_discount']=configured_profit/configured_price
        rec['activity_price_at_3pct_margin_with_extra_giftbox_13']=max((cost+13+59.5)/.87,(cost+13+14.5)/.65)
    rec['activity_cached_status']='valid' if isinstance(p,(int,float)) and p>0 else 'unavailable activity price: missing/error Q; T may fallback to normal gross margin O'
    if isinstance(p,(int,float)) and isinstance(cost,(int,float)) and p>0:
        profit=p-max(p*.22,45)-14.5-p*.10-cost
        rec['activity_profit_independent']=profit;rec['activity_margin_independent']=profit/p
        if abs(profit-val(cells,'S'))>1e-7:discrepancies.append({'row':n,'article_no':rec['article_no'],'recomputed':profit,'cached':val(cells,'S')})
    records.append(rec)
index={'source':str(src),'source_sheet':'Sheet1','source_range':'A1:T98','record_count':len(records),'id_note':'A列仅标注ID；50条有效数字、47条为#N/A，需与后台商品ID交叉核验；无款式名或条码字段，C列是图片。货号索引已统一大写。','formula_notes':{'current':'N=F-G-MAX(F*22%,45)-I-J-K-L-M, O=N/F; I=5,J=3,K=5,L=3,M=F*3%。','activity':'R=D*Q*0.1; S=R-MAX(R*22%,45)-5-3-5-1.5-R*10%-G; T=IF(R>0,S/R,O)。','activity_deductions':'平台扣点22%，最低45元；包装5、物料3、运费5、操作费1.5；另扣活动价10%（公式未标注具体业务含义，不能擅自认定税费）；产品成本G。','giftbox':'工作簿未设礼盒增量成本/规格字段。用户本轮确认两个升级礼盒成本都是13元；按产品成本G之外增加13计算，并保留公式原包装5元。品牌礼盒未另加13。','discount_units':'E/Q使用7表示7折、8表示8折；活动价按照出价D×折扣/10。','three_pct_threshold':'在所有S公式均一致且无额外活动费用时，3%毛利临界价=max((成本+59.5)/0.87,(成本+14.5)/0.65)；超过临界价才高于3%。','cache':'66条有效活动价的S/T缓存经独立复算一致；Q68:Q97空白使R=0、T回退普通毛利率O，Q98外链缓存#N/A；不能直接采用这些T判断活动。24条G13:G36成本来自外链缓存。原文件未改写。','break_even_P':'P是普通毛利口径的保本估算，不能作为活动3%门槛；活动应重新计算。'},'missing_cost_price_discount_rows':missing,'invalid_product_id_rows':[r['row'] for r in records if not r['product_id'].isdigit()],'external_cost_cache_rows':[r['row'] for r in records if 'G' in r['formulas']],'activity_recalculation_discrepancies':discrepancies,'formula_groups':[{'column':col,'formula_pattern':f,'rows':rs} for (col,f),rs in group.items()],'records':records,'by_article_no':{r['article_no'].upper():r for r in records},'by_product_id':{r['product_id']:r for r in records if r['product_id'].isdigit()}}
pathlib.Path('work/catalog_costs.json').write_text(json.dumps(index,ensure_ascii=False,indent=2,default=str))
print(json.dumps({'record_count':len(records),'missing_rows':missing,'valid_product_ids':len(index['by_product_id']),'invalid_product_id_count':len(index['invalid_product_id_rows']),'external_cost_cache_count':len(index['external_cost_cache_rows']),'duplicates_article_no':[k for k,v in collections.Counter(r['article_no'].upper() for r in records).items() if v>1],'discrepancies':discrepancies,'activity_margin_at_configured_below_3pct':[{'row':r['row'],'article_no':r['article_no'],'margin':r['activity_margin_at_configured_discount']} for r in records if r['activity_margin_at_configured_discount']<.03]},ensure_ascii=False,default=str))
